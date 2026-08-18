import sys
import types

from docx import Document

from app.db.schemas import TemplateField, TemplateSection
from app.services import schema_inference
from app.services.extraction_service import ExtractionService
from app.services.model.qwen_vl import QwenVLProvider
from app.services.schema_inference import (
    infer_schema_sections_with_page_count,
    extract_deterministic_page_data,
    compare_page_to_source,
)
from app.services.page_analysis_service import build_page_text_from_fields


def test_extract_deterministic_page_data_and_compare_to_source():
    source_text = "Project Name: Example Project\nDocument Number: DOC-100\n\nTable:\nA | B\n1 | 2"
    page_data = extract_deterministic_page_data(source_text)

    assert page_data["table_count"] >= 1
    assert "Project Name" in page_data["field_labels"]
    assert page_data["layout_score"] >= 0.0

    comparison = compare_page_to_source(source_text, page_data["normalized_text"])
    assert comparison["coverage_ratio"] > 0.0
    assert comparison["table_match_ratio"] >= 0.0
    assert "page_coverage" in comparison


def test_resolve_soffice_uses_configured_executable(monkeypatch, tmp_path):
    configured = tmp_path / "soffice.exe"
    configured.write_text("stub")
    monkeypatch.setattr(
        schema_inference,
        "settings",
        types.SimpleNamespace(LIBREOFFICE_PATH=str(configured)),
    )
    monkeypatch.setattr(schema_inference.shutil, "which", lambda cmd: None)

    assert schema_inference._resolve_soffice_executable() == str(configured)


def test_resolve_soffice_prefers_executable_over_com_launcher(monkeypatch):
    monkeypatch.setattr(
        schema_inference,
        "settings",
        types.SimpleNamespace(LIBREOFFICE_PATH=None),
    )

    def fake_which(command):
        if command == "soffice.exe":
            return r"C:\Program Files\LibreOffice\program\soffice.exe"
        if command == "soffice":
            return r"C:\Program Files\LibreOffice\program\soffice.COM"
        if command == "soffice.com":
            return r"C:\Program Files\LibreOffice\program\soffice.com"
        return None

    monkeypatch.setattr(schema_inference.shutil, "which", fake_which)

    assert schema_inference._resolve_soffice_executable() == r"C:\Program Files\LibreOffice\program\soffice.exe"


def test_word_conversion_initializes_com_before_dispatch(monkeypatch, tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"dummy")
    calls = []

    class FakeWordDocument:
        def SaveAs2(self, target_path, FileFormat):
            calls.append(("saveas", target_path, FileFormat))

        def Close(self, SaveChanges):
            calls.append(("close", SaveChanges))

    class FakeDocuments:
        @staticmethod
        def Open(path, ReadOnly=True):
            calls.append(("open", path, ReadOnly))
            return FakeWordDocument()

    class FakeWordApplication:
        Visible = False
        Documents = FakeDocuments()

        def Quit(self):
            calls.append(("quit",))

    pythoncom = types.ModuleType("pythoncom")
    pythoncom.CoInitialize = lambda *args, **kwargs: calls.append(("coinitialize", args, kwargs))
    pythoncom.CoUninitialize = lambda *args, **kwargs: calls.append(("couninitialize", args, kwargs))
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)

    fake_win32com = types.ModuleType("win32com")
    fake_client = types.ModuleType("win32com.client")

    def fake_dispatch(app_name):
        calls.append(("dispatch", app_name))
        return FakeWordApplication()

    fake_client.Dispatch = fake_dispatch
    fake_win32com.client = fake_client
    monkeypatch.setitem(sys.modules, "win32com", fake_win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", fake_client)
    monkeypatch.setattr(schema_inference.shutil, "which", lambda cmd: "winword" if cmd == "winword" else None)

    result = schema_inference._convert_doc_to_docx(source)

    assert result == tmp_path / "legacy.docx"
    assert ("coinitialize", (), {}) in calls
    assert calls.index(("coinitialize", (), {})) < calls.index(("dispatch", "Word.Application"))


def test_infer_schema_sections_extracts_table_fields_and_page_count(tmp_path):
    doc = Document()
    doc.add_heading("Project Details")
    doc.add_paragraph("Project Name: Example Project")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Document Number"
    table.rows[0].cells[1].text = ""
    table.rows[1].cells[0].text = "Prepared By"
    table.rows[1].cells[1].text = ""

    path = tmp_path / "sample.docx"
    doc.save(path)

    sections, text_preview, page_count, preview_html, page_images, page_html = infer_schema_sections_with_page_count(path)

    labels = [field["field_label"] for section in sections for field in section["fields"]]
    assert "Project Name" in labels
    assert "Document Number" in labels
    assert "Prepared By" in labels
    assert "Example Project" in text_preview
    assert page_count >= 1
    assert "Project Details" in preview_html
    assert "Document Number" in preview_html


def test_schema_inference_ignores_empty_table_cells():
    class FakeCell:
        def __init__(self, text=""):
            self.text = text
            self.paragraphs = []
            self.tables = []

    class FakeRow:
        def __init__(self, cells):
            self.cells = cells

    class FakeTable:
        def __init__(self, rows):
            self.rows = rows

    row = FakeRow([None, FakeCell("Project Name"), FakeCell("")])
    table = FakeTable([row])

    assert schema_inference._extract_row_cells(row) == ["Project Name", ""]
    assert schema_inference._render_table_plain(table) == ""
    assert "Project Name" in schema_inference._render_docx_preview_html(FakeTable([FakeRow([FakeCell("Project Name")])]))


def test_infer_schema_sections_extracts_ups_form_table_values(tmp_path):
    doc = Document()
    doc.add_heading("4.8 UN-INTERRUPTED POWER SUPPLY (UPS) SYSTEM")

    table = doc.add_table(rows=1, cols=2)
    rows = [
        ["a) AC Input supply", "415V ±10 %, 50Hz, 3ph, 4Wire AC"],
        ["b) Output voltage & variation", "110 V ±2 % - For Instrument"],
        ["UPS", ""],
        ["c) Output frequency & variation", "50 Hz ±1 %"],
        ["d) Rated KVA", ""],
        ["Instrument UPS", "Later"],
        ["Power UPS", "Later (Phase I Critical loads.)"],
        ["e) Static Transfer switch", "Required"],
        ["f) Audio visual alarm", "Required"],
        ["g) AC distribution board", "Required"],
        ["h) type", "Lead acid Sealed low maintenance"],
        ["i) Duration of back-up for rated output", ""],
        ["- Power UPS", "6 Minutes"],
        ["- Instrument UPS", "30 Minutes"],
    ]
    for index, values in enumerate(rows):
        if index >= len(table.rows):
            table.add_row()
        row = table.rows[index]
        for col_index, value in enumerate(values):
            row.cells[col_index].text = value

    path = tmp_path / "ups_form.docx"
    doc.save(path)

    sections, _, _, _, _, _ = infer_schema_sections_with_page_count(path)
    labels = [field["field_label"] for section in sections for field in section["fields"]]

    assert "AC Input supply" in labels
    assert "Output voltage & variation" in labels
    assert "Output frequency & variation" in labels
    assert "Static Transfer switch" in labels
    assert "Audio visual alarm" in labels
    assert "Instrument UPS" in labels
    assert "Power UPS" in labels
    assert "Duration of back-up for rated output" in labels


def test_build_page_text_from_fields_keeps_labels_and_values():
    fields = [
        {"field_label": "Project Name", "value": "Example Project"},
        {"field_label": "Prepared By", "value": "Jane Smith"},
    ]

    text = build_page_text_from_fields(fields)

    assert "Project Name: Example Project" in text
    assert "Prepared By: Jane Smith" in text


def test_extract_batches_large_sections_into_small_prompts(monkeypatch):
    provider = QwenVLProvider()
    provider.base_url = "http://localhost:11434"
    observed_batches = []

    def fake_extract_single_schema(schema, pages):
        observed_batches.append(len(schema["sections"][0]["fields"]))
        return {
            "template_id": schema["template_id"],
            "template_version": "1.0",
            "fields": [
                {
                    "field_id": field["field_id"],
                    "field_name": field["field_label"],
                    "value": "example",
                    "confidence": 1.0,
                    "source": {"page_number": 1, "source_text": "example"},
                }
                for field in schema["sections"][0]["fields"]
            ],
        }

    monkeypatch.setattr(provider, "_extract_single_schema", fake_extract_single_schema)

    schema = {
        "template_id": "demo",
        "version": "1.0",
        "sections": [
            {
                "section_name": "Large section",
                "fields": [
                    {"field_id": f"field_{i}", "field_label": f"Field {i}", "extraction_hint": "value"}
                    for i in range(45)
                ],
            }
        ],
    }

    result = provider.extract(schema, [{"page_number": 1, "text": "Example text"}])

    assert len(result["fields"]) == 45
    assert observed_batches
    assert max(observed_batches) <= provider.MAX_FIELDS_PER_PROMPT
    assert sum(observed_batches) == 45


def test_template_schema_models_preserve_page_numbers():
    field = TemplateField(
        field_id="project_name",
        field_label="Project Name",
        data_type="string",
        required=True,
        page_number=2,
    )
    section = TemplateSection(
        section_id="header",
        section_name="Header",
        page_number=2,
        fields=[field],
    )

    assert field.page_number == 2
    assert section.page_number == 2
    assert section.fields[0].page_number == 2


def test_settings_loads_local_env_override_for_ollama_url():
    config_module = __import__("app.core.config", fromlist=["Settings"])
    env_file_cfg = config_module.Settings.model_config.get("env_file")

    assert env_file_cfg is not None
    env_files = env_file_cfg if isinstance(env_file_cfg, (list, tuple)) else [env_file_cfg]
    normalized = [str(path).replace("\\", "/") for path in env_files]
    assert any("backend/.env" in value for value in normalized)


def test_qwen_vl_candidate_urls_include_docker_and_host_fallbacks():
    provider = QwenVLProvider()
    provider.base_url = "http://localhost:11434"

    urls = provider._candidate_urls()

    assert "http://localhost:11434" in urls
    assert "http://localhost:11435" in urls
    assert "http://host.docker.internal:11434" in urls
    assert "http://ollama:11434" in urls


def test_validate_output_rejects_missing_required_fields():
    service = ExtractionService.__new__(ExtractionService)
    schema = {
        "template_id": "tpl-1",
        "version": "v1",
        "sections": [
            {
                "section_name": "Header",
                "fields": [
                    {"field_id": "project_name", "field_label": "Project Name", "required": True},
                    {"field_id": "project_id", "field_label": "Project ID", "required": True},
                ],
            }
        ],
    }
    model_output = {
        "template_id": "tpl-1",
        "template_version": "v1",
        "fields": [
            {"field_id": "project_name", "field_name": "Project Name", "value": "PASHMINA PROJECT", "confidence": 0.98, "source": {"page_number": 1, "source_text": "PASHMINA PROJECT"}},
            {"field_id": "project_id", "field_name": "Project ID", "value": None, "confidence": 0.0, "source": None},
        ],
    }

    is_valid, message = service._validate_output(model_output, schema)

    assert is_valid is False
    assert "Project ID" in str(message)
