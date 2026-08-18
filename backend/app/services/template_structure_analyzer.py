"""
Template Structure Analysis and Signature Generation

This module analyzes master template documents to extract structural
characteristics (page count, table structure, paragraph layout, etc.)
and generates a structure signature used to validate that generated
documents maintain the same layout as the master template.

Core Principle: Master template structure is IMMUTABLE - only field
values can change between input and output.
"""
import hashlib
import json
import logging
from pathlib import Path
from typing import Any, Optional
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

try:
    import fitz
except ImportError:
    fitz = None

logger = logging.getLogger(__name__)


class TableStructureInfo:
    """Captures the structural properties of a table."""
    
    def __init__(self, table):
        self.row_count = len(table.rows)
        self.col_count = len(table.columns)
        # Store cell text as structural fingerprint
        self.cell_structure = []
        for row_idx, row in enumerate(table.rows):
            row_data = []
            for col_idx, cell in enumerate(row.cells):
                # Only capture structural info, not values
                cell_text = cell.text.strip()
                is_empty = len(cell_text) == 0
                row_data.append({
                    "row": row_idx,
                    "col": col_idx,
                    "is_empty": is_empty,
                    "para_count": len(cell.paragraphs),
                })
            self.cell_structure.append(row_data)
    
    def to_dict(self) -> dict[str, Any]:
        """Convert to serializable dict."""
        return {
            "row_count": self.row_count,
            "col_count": self.col_count,
            "cell_structure": self.cell_structure,
        }
    
    def signature_hash(self) -> str:
        """Generate hash of table structure for comparison."""
        sig_str = json.dumps({
            "rows": self.row_count,
            "cols": self.col_count,
            "cells": len(self.cell_structure) * len(self.cell_structure[0]) if self.cell_structure else 0,
        }, sort_keys=True)
        return hashlib.sha256(sig_str.encode()).hexdigest()[:16]


class ParagraphInfo:
    """Captures structural properties of paragraphs."""
    
    def __init__(self, para: Paragraph):
        self.text = para.text.strip()
        self.style = para.style.name if para.style else "Normal"
        self.level = self._get_outline_level(para)
        self.alignment = str(para.alignment) if para.alignment is not None else "None"
        self.is_empty = len(self.text) == 0
        self.run_count = len(para.runs)

    @staticmethod
    def _get_outline_level(para: Paragraph) -> int | None:
        """Read the paragraph's outline level directly from its XML.

        python-docx's ParagraphFormat has no `outline_level` property (this
        was called unconditionally in every ParagraphInfo, so
        analyze_docx_structure raised AttributeError on every single
        document, every time - the integrity check that's supposed to
        guarantee the output stays structurally identical to the master
        template was never actually completing). <w:outlineLvl> lives on
        pPr directly, so read it from there instead.
        """
        try:
            pPr = para._p.find(qn("w:pPr"))
            if pPr is None:
                return None
            outline_el = pPr.find(qn("w:outlineLvl"))
            if outline_el is None:
                return None
            val = outline_el.get(qn("w:val"))
            return int(val) if val is not None else None
        except (AttributeError, ValueError, TypeError):
            return None
    
    def to_dict(self) -> dict[str, Any]:
        """Convert to serializable dict."""
        return {
            "text_length": len(self.text),
            "style": self.style,
            "level": self.level,
            "alignment": self.alignment,
            "is_empty": self.is_empty,
            "run_count": self.run_count,
        }
    
    def signature_hash(self) -> str:
        """Generate hash of paragraph structure (NOT content)."""
        sig_str = json.dumps({
            "style": self.style,
            "level": self.level,
            "alignment": self.alignment,
            "run_count": self.run_count,
        }, sort_keys=True)
        return hashlib.sha256(sig_str.encode()).hexdigest()[:16]


class DocumentStructureSignature:
    """Complete structural fingerprint of a document."""
    
    def __init__(self):
        self.page_count = 0
        self.table_count = 0
        self.paragraph_count = 0
        self.has_headers = False
        self.has_footers = False
        self.has_sections = False
        # Sequence of block types for overall structure
        self.block_sequence = []  # ["para", "table", "para", ...]
        # Hashes for comparison
        self.structure_hashes = []  # List of hashes for each major block
        self.overall_hash = ""
    
    def to_dict(self) -> dict[str, Any]:
        """Convert to serializable dict."""
        return {
            "page_count": self.page_count,
            "table_count": self.table_count,
            "paragraph_count": self.paragraph_count,
            "has_headers": self.has_headers,
            "has_footers": self.has_footers,
            "has_sections": self.has_sections,
            "block_sequence": self.block_sequence,
            "structure_hashes": self.structure_hashes,
            "overall_hash": self.overall_hash,
        }
    
    @classmethod
    def from_dict(cls, data: dict) -> "DocumentStructureSignature":
        """Create signature from dict."""
        sig = cls()
        sig.page_count = data.get("page_count", 0)
        sig.table_count = data.get("table_count", 0)
        sig.paragraph_count = data.get("paragraph_count", 0)
        sig.has_headers = data.get("has_headers", False)
        sig.has_footers = data.get("has_footers", False)
        sig.has_sections = data.get("has_sections", False)
        sig.block_sequence = data.get("block_sequence", [])
        sig.structure_hashes = data.get("structure_hashes", [])
        sig.overall_hash = data.get("overall_hash", "")
        return sig


class TemplateStructureAnalyzer:
    """Analyzes master templates to capture structural characteristics."""
    
    @staticmethod
    def analyze_docx_structure(docx_path: Path) -> DocumentStructureSignature:
        """
        Analyze a DOCX file and generate a structure signature.
        
        This captures all structural elements but NOT content/values,
        so it can be used to validate that generated documents maintain
        the same layout as the master template.
        """
        sig = DocumentStructureSignature()
        
        try:
            doc = DocxDocument(str(docx_path))
        except Exception as e:
            logger.error(f"Failed to open DOCX: {e}")
            return sig
        
        # Analyze document structure
        sig.has_sections = len(doc.sections) > 0
        sig.has_headers = any(
            section.header.paragraphs for section in doc.sections
        )
        sig.has_footers = any(
            section.footer.paragraphs for section in doc.sections
        )
        
        # Count pages by looking at page breaks
        page_breaks = 0
        block_sequence = []
        structure_hashes = []
        
        # Walk document blocks in order
        body = doc.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):  # Paragraph
                para = Paragraph(child, doc)
                sig.paragraph_count += 1
                block_sequence.append("para")
                
                # Check for page break in this paragraph
                for run in para.runs:
                    if run.element.xpath('.//w:br[@w:type="page"]'):
                        page_breaks += 1
                
                # Add structural hash (not content). A single malformed
                # paragraph shouldn't take down analysis of the entire
                # document - record it as an empty-signature paragraph
                # rather than raising.
                try:
                    para_info = ParagraphInfo(para)
                    structure_hashes.append(para_info.signature_hash())
                except Exception as e:
                    logger.warning("Failed to analyze paragraph structure: %s", e)
                    structure_hashes.append("")
                
            elif child.tag == qn("w:tbl"):  # Table
                table = Table(child, doc)
                sig.table_count += 1
                block_sequence.append("table")
                
                # Add structural hash
                table_info = TableStructureInfo(table)
                structure_hashes.append(table_info.signature_hash())
        
        # Page count is breaks + 1 (every document has at least 1 page)
        sig.page_count = page_breaks + 1
        sig.block_sequence = block_sequence
        sig.structure_hashes = structure_hashes
        
        # Generate overall hash
        sig.overall_hash = TemplateStructureAnalyzer._compute_overall_hash(sig)
        
        return sig
    
    @staticmethod
    def analyze_pdf_structure(pdf_path: Path) -> DocumentStructureSignature:
        """
        Analyze a PDF file and generate a structure signature.
        
        PDF analysis is more limited than DOCX since we can't access
        the underlying structure - we primarily count pages.
        """
        sig = DocumentStructureSignature()
        
        if fitz is None:
            logger.warning("PyMuPDF (fitz) not available for PDF analysis")
            return sig
        
        try:
            pdf_doc = fitz.open(str(pdf_path))
            sig.page_count = len(pdf_doc)
            pdf_doc.close()
        except Exception as e:
            logger.error(f"Failed to analyze PDF: {e}")
        
        return sig
    
    @staticmethod
    def _compute_overall_hash(sig: DocumentStructureSignature) -> str:
        """Compute overall structure hash for the entire document."""
        sig_data = {
            "pages": sig.page_count,
            "tables": sig.table_count,
            "paragraphs": sig.paragraph_count,
            "headers": sig.has_headers,
            "footers": sig.has_footers,
            "sections": sig.has_sections,
            "block_sequence_length": len(sig.block_sequence),
            "hashes_count": len(sig.structure_hashes),
        }
        sig_str = json.dumps(sig_data, sort_keys=True)
        return hashlib.sha256(sig_str.encode()).hexdigest()[:16]
    
    @staticmethod
    def compare_signatures(
        original: DocumentStructureSignature,
        generated: DocumentStructureSignature,
    ) -> tuple[bool, list[str]]:
        """
        Compare two structure signatures.
        
        Returns: (is_identical, list_of_differences)
        
        For template population, structural differences indicate a failed
        replacement - only values should change, not structure.
        """
        differences = []
        
        if original.page_count != generated.page_count:
            differences.append(
                f"Page count mismatch: {original.page_count} vs {generated.page_count}"
            )
        
        if original.table_count != generated.table_count:
            differences.append(
                f"Table count mismatch: {original.table_count} vs {generated.table_count}"
            )
        
        if original.paragraph_count != generated.paragraph_count:
            # Allow some variance in paragraph count due to reformatting
            if abs(original.paragraph_count - generated.paragraph_count) > 5:
                differences.append(
                    f"Paragraph count variance: {original.paragraph_count} vs {generated.paragraph_count}"
                )
        
        if original.has_headers != generated.has_headers:
            differences.append(
                f"Header presence mismatch: {original.has_headers} vs {generated.has_headers}"
            )
        
        if original.has_footers != generated.has_footers:
            differences.append(
                f"Footer presence mismatch: {original.has_footers} vs {generated.has_footers}"
            )
        
        if original.block_sequence != generated.block_sequence:
            differences.append(
                "Block sequence mismatch (layout structure changed)"
            )
        
        is_identical = len(differences) == 0
        return is_identical, differences