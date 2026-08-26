import os
import re
import sys
import shutil
import subprocess
from pathlib import Path
from uuid import uuid4
from datetime import datetime
from typing import List, cast, Optional, Any
from sqlalchemy.orm import Session
from docx import Document as DocxDocument
from PIL import Image
import fitz
import pytesseract
from app.core.config import settings
from app.db.models import Document, DocumentPage, Project

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".png", ".jpg", ".jpeg", ".tiff", ".bmp"}

# Characters a genuine text layer is overwhelmingly made of: letters,
# digits, whitespace, and common punctuation/symbols found in engineering
# specs (units, dashes, slashes, etc).
_PLAUSIBLE_TEXT_CHARS = re.compile(r"[A-Za-z0-9\s\.,;:\-\+\/\(\)%'\"&#@=_]")

# A broken font's ToUnicode mapping usually decodes a small handful of
# glyph codes wrong and repeats them - so the telltale sign in practice is
# a short chunk (roughly one glyph-cluster's worth of characters) repeating
# immediately several times in a row, e.g. "1545-E0..x1545-E0..x1545-E0..x".
# This catches that case even when the wrong characters happen to be
# digits/dashes that would otherwise look "plausible" by charset alone.
_REPEATED_CHUNK = re.compile(r"(.{3,25}?)\1{2,}")
# Decorative table/section rules ("----", "....", "____") are legitimate
# and repeat by design - don't let those alone trip the repetition check.
_BORDER_CHARS = set("-_.=*# ")


def _looks_like_garbage_text(text: str, min_ratio: float = 0.6) -> bool:
    """PyMuPDF's page.get_text() returns whatever the PDF's font encoding
    maps character codes to. Some PDFs (commonly CAD/drafting-tool exports)
    embed subset fonts with a broken or missing ToUnicode CMap - get_text()
    still returns a non-empty string, but it's garbled glyph soup, not real
    text (this is what produces repeating symbol strings like
    "\\1545-E0(check)x1545-" instead of the actual field values).
    Treating "non-empty" as "usable" let that garbage flow straight into
    the extraction model instead of falling back to OCR. Two independent
    checks: an unusual-character ratio (catches non-printable glyph soup),
    and immediate short-chunk repetition covering a large share of the
    text (catches broken mappings that happen to decode to plausible
    characters like digits/dashes, which the ratio check alone misses).
    """
    stripped = text.strip()
    if not stripped:
        return True
    plausible = len(_PLAUSIBLE_TEXT_CHARS.findall(stripped))
    if (plausible / len(stripped)) < min_ratio:
        return True
    m = _REPEATED_CHUNK.search(stripped)
    if m:
        chunk = m.group(1).strip()
        is_border_style = set(chunk) <= _BORDER_CHARS
        if not is_border_style and len(chunk) >= 3 and len(m.group(0)) >= 0.3 * len(stripped):
            return True
    return False

class DocumentService:
    def __init__(self, db: Session):
        self.db = db
        self.originals_path = Path(settings.STORAGE_ORIGINALS_PATH)
        self.processed_path = Path(settings.STORAGE_PROCESSED_PATH)
        self.pages_path = Path(settings.STORAGE_PAGES_PATH)
        self.originals_path.mkdir(parents=True, exist_ok=True)
        self.processed_path.mkdir(parents=True, exist_ok=True)
        self.pages_path.mkdir(parents=True, exist_ok=True)

    def _validate_extension(self, file_name: str) -> str:
        ext = Path(file_name).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise ValueError("Unsupported file type")
        return ext

    def save_document(self, project_id: int, upload_file) -> Document:
        original_filename = upload_file.filename
        extension = self._validate_extension(original_filename)
        stored_filename = f"{uuid4().hex}{extension}"
        destination = self.originals_path / stored_filename

        total_size = 0
        with destination.open("wb") as out_file:
            while True:
                chunk = upload_file.file.read(1024 * 1024)
                if not chunk:
                    break
                total_size += len(chunk)
                if total_size > settings.MAX_UPLOAD_SIZE:
                    out_file.close()
                    destination.unlink(missing_ok=True)
                    raise ValueError("File too large")
                out_file.write(chunk)

        document = Document(
            project_id=project_id,
            original_filename=original_filename,
            stored_filename=stored_filename,
            file_path=str(destination),
            file_type=extension.lstrip("."),
            file_size=total_size,
            upload_status="uploaded",
        )
        self.db.add(document)
        self.db.commit()
        self.db.refresh(document)

        self.process_document(document)
        return document

    def list_documents(self, project_id: int) -> List[Document]:
        return self.db.query(Document).filter_by(project_id=project_id).order_by(Document.created_at.desc()).all()

    def to_response(self, document: Document) -> dict:
        created_at = cast(datetime | None, document.created_at)
        return {
            "id": document.id,
            "project_id": document.project_id,
            "original_filename": document.original_filename,
            "stored_filename": document.stored_filename,
            "file_type": document.file_type,
            "file_size": document.file_size,
            "page_count": document.page_count,
            "upload_status": document.upload_status,
            "created_at": created_at.isoformat() if created_at else None,
        }

    def get_document(self, document_id: int) -> Document | None:
        return self.db.query(Document).filter_by(id=document_id).first()

    def get_page(self, document_id: int, page_number: int) -> Optional[DocumentPage]:
        return (
            self.db.query(DocumentPage)
            .filter_by(document_id=document_id, page_number=page_number)
            .first()
        )

    def process_document(self, document: Document) -> None:
        stored_filename = cast(str, document.stored_filename)
        file_type = cast(str, document.file_type)

        document_path = self.originals_path / stored_filename
        if not document_path.exists():
            setattr(document, "upload_status", "failed")
            self.db.add(document)
            self.db.commit()
            return

        pages = []
        if file_type == "pdf":
            pages = self._extract_pdf_pages(document, document_path)
        elif file_type in {"docx", "doc"}:
            pages = self._extract_docx_pages(document, document_path)
        elif file_type in {"png", "jpg", "jpeg", "tiff", "bmp"}:
            pages = self._extract_image_pages(document, document_path)
        else:
            pages = []

        setattr(document, "page_count", len(pages))
        setattr(document, "upload_status", "processed")
        self.db.add(document)
        self.db.commit()

        doc_id = cast(int, document.id)
        for page_number, page_text, image_path in pages:
            page = DocumentPage(
                document_id=doc_id,
                page_number=page_number,
                text=page_text,
                image_path=image_path,
            )
            self.db.add(page)
        self.db.commit()

    def _extract_pdf_pages(self, document: Document, path: Path):
        pages = []
        doc_id = cast(int, document.id)
        try:
            pdf_doc = cast(Any, getattr(fitz, "open")(str(path)))
            for number, page in enumerate(pdf_doc, start=1):
                text = page.get_text() or ""
                image_path = None
                if _looks_like_garbage_text(text):
                    pix = page.get_pixmap()
                    image_path = self._save_page_image(doc_id, number, pix.pil_tobytes(format="PNG"))
                    text = self._ocr_image_bytes(pix.tobytes("png"))
                else:
                    if settings.STORAGE_PAGES_PATH:
                        image_path = self._save_page_image(doc_id, number, page.get_pixmap().tobytes("png"))
                pages.append((number, text, image_path))
        except Exception:
            pass
        return pages

    def _extract_docx_pages(self, document: Document, path: Path):
        doc_id = cast(int, document.id)
        target_path = path
        if path.suffix.lower() == ".doc":
            target_path = self._convert_doc_to_docx(path)

        from app.services.schema_inference import _docx_to_pdf, _PageLocator, _iter_block_items
        from docx.table import Table as _DocxTable
        import docx as _docx_module

        table_rows_by_page: dict[int, list[str]] = {}
        try:
            structure_doc = _docx_module.Document(str(target_path))
        except Exception:
            structure_doc = None

        pdf_path = None
        try:
            pdf_path = _docx_to_pdf(target_path)
            pdf_doc = cast(Any, getattr(fitz, "open")(str(pdf_path)))
            page_texts = [pdf_doc[i].get_text() or "" for i in range(len(pdf_doc))]

            if structure_doc is not None and page_texts:
                locator = _PageLocator(page_texts)
                current_page = 1
                for block in _iter_block_items(structure_doc):
                    if isinstance(block, _DocxTable):
                        for row in block.rows:
                            try:
                                cells = [c.text.strip() for c in row.cells]
                            except (IndexError, AttributeError):
                                continue
                            deduped: list[str] = []
                            for c in cells:
                                if not deduped or deduped[-1] != c:
                                    deduped.append(c)
                            row_text = "\t".join(c for c in deduped if c)
                            if row_text:
                                page_for_row = locator.locate(row_text, current_page)
                                current_page = page_for_row
                                table_rows_by_page.setdefault(page_for_row, []).append(row_text)
                        continue
                    text = block.text.strip()
                    if text:
                        current_page = locator.locate(text, current_page)

            pages = []
            for page_index in range(len(pdf_doc)):
                page = pdf_doc[page_index]
                number = page_index + 1
                text_layer = page_texts[page_index]

                mat = fitz.Matrix(150 / 72, 150 / 72)
                pix = page.get_pixmap(matrix=mat, alpha=False)
                image_bytes = pix.tobytes("png")
                try:
                    ocr_text = self._ocr_image_bytes(image_bytes)
                except Exception:
                    ocr_text = ""

                combined_text = text_layer
                if ocr_text.strip() and ocr_text.strip() not in text_layer:
                    combined_text = f"{text_layer}\n{ocr_text}".strip()

                table_rows = table_rows_by_page.get(number)
                if table_rows:
                    combined_text = (
                        f"{combined_text}\n\nSTRUCTURED TABLE DATA (columns "
                        f"tab-separated, in source order):\n" + "\n".join(table_rows)
                    ).strip()

                image_path = self._save_page_image(doc_id, number, image_bytes)
                pages.append((number, combined_text, image_path))
            pdf_doc.close()
            if pages:
                return pages
        except Exception:
            pass
        finally:
            if pdf_path is not None:
                try:
                    pdf_path.unlink()
                except Exception:
                    pass

        try:
            from app.services.schema_inference import extract_docx_text_pages

            text_pages = extract_docx_text_pages(target_path)
            if text_pages:
                return [(num, text, None) for num, text in text_pages]
        except Exception:
            pass

        try:
            doc = DocxDocument(str(target_path))
            text = "\n".join(p.text for p in doc.paragraphs if p.text)
            return [(1, text, None)]
        except Exception:
            return []

    def _extract_image_pages(self, document: Document, path: Path):
        doc_id = cast(int, document.id)
        try:
            text = self._ocr_image(path)
            image_path = self._save_page_image(doc_id, 1, path.read_bytes())
            return [(1, text, image_path)]
        except Exception:
            return []

    def _convert_doc_to_docx(self, path: Path) -> Path:
        soffice_executable = None
        configured = getattr(settings, "LIBREOFFICE_PATH", None)

        # 1. First check if configured path directly exists
        if configured and Path(configured).expanduser().exists():
            soffice_executable = str(Path(configured).expanduser())
        else:
            # 2. Dynamic binary resolution prioritizing Linux/standard executable names
            soffice_executable = (
                shutil.which("libreoffice")
                or shutil.which("soffice")
                or shutil.which("soffice.exe")
                or shutil.which("soffice.com")
            )

        if soffice_executable is not None:
            output_dir = self.processed_path
            output_dir.mkdir(parents=True, exist_ok=True)
            command = [
                soffice_executable,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(output_dir),
                str(path),
            ]
            subprocess.run(command, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            converted = output_dir / (path.stem + ".docx")
            return converted

        # Word COM Automation (Windows-only fallback)
        if sys.platform == "win32" and (
            shutil.which("winword") is not None
            or Path("C:/Program Files/Microsoft Office/root/Office16/WINWORD.EXE").exists()
        ):
            try:
                import pythoncom  # type: ignore[import-not-found, import-untyped]
                import win32com.client  # type: ignore[import-not-found, import-untyped]
            except ImportError as exc:  # pragma: no cover - Windows-only dependency path.
                raise FileNotFoundError(
                    "Neither LibreOffice nor Word automation is available for .doc conversion. "
                    "Please convert the file to .docx before uploading."
                ) from exc

            output_dir = self.processed_path
            output_dir.mkdir(parents=True, exist_ok=True)
            output_path = output_dir / (path.stem + ".docx")
            pythoncom.CoInitialize()
            word = None
            try:
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(str(path), ReadOnly=True)
                doc.SaveAs2(str(output_path), FileFormat=16)
                doc.Close(SaveChanges=False)
            finally:
                if word is not None:
                    try:
                        word.Quit()
                    except Exception:
                        pass
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass
            return output_path

        raise FileNotFoundError(
            "Neither LibreOffice 'soffice' nor Microsoft Word is available to convert .doc templates to .docx. "
            "Please convert the file to .docx before uploading."
        )

    def _ocr_image(self, path: Path) -> str:
        image = Image.open(path)
        return pytesseract.image_to_string(image)

    def _ocr_image_bytes(self, image_bytes: bytes) -> str:
        from io import BytesIO

        image = Image.open(BytesIO(image_bytes)).convert("RGB")
        return pytesseract.image_to_string(image)

    def _save_page_image(self, document_id: int, page_number: int, image_bytes: bytes) -> str:
        page_dir = self.pages_path / str(document_id)
        page_dir.mkdir(parents=True, exist_ok=True)
        image_path = page_dir / f"page_{page_number}.png"
        image_path.write_bytes(image_bytes)
        return str(image_path.resolve())